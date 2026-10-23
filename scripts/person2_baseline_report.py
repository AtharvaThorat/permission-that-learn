"""Build the Person 2 unprotected-baseline report from experiment run folders.

Reads a registry written by `person2_unprotected_baseline.ipynb` (one entry per model, with its
run folder or the reason it has none) and writes, without modifying any run folder:

  unprotected_baseline_results.xlsx   sheets: Summary, Episodes, Run info
  summary.csv, episodes.csv           the same tables as CSV
  charts/1_task_utility.png           utility per model
  charts/2_model_activity.png         actor requests, tool calls proposed and executed
  charts/3_run_cost.png               tokens and episode duration per model

Metrics that a run did not produce stay empty; they are never replaced with zero.
Needs only the standard library, openpyxl and matplotlib (both preinstalled on Colab).

  python scripts/person2_baseline_report.py --registry <dir>/registry.json --output <dir>/report
"""

import argparse
import csv
import json
from pathlib import Path

# Colorblind-safe categorical palette, assigned by registry order so a model keeps its color.
SERIES_COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]
SURFACE, TEXT, MUTED_TEXT, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e4e3df"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def text_of(content) -> str:
    """AgentDojo message content is a list of {'type': 'text', 'content': ...} blocks, or null."""
    if not content:
        return ""
    return "".join(block.get("content") or "" for block in content)


def episode_rows(entry: dict) -> list[dict]:
    """One row per episode attempt in the entry's run folder."""
    rows = []
    run_dir = Path(entry["run_dir"])
    for result_path in sorted(run_dir.glob("episodes/*/*/result.json")):
        result = read_json(result_path)
        conversation_path = result_path.parent / result["artifacts"]["conversation"]
        messages = read_json(conversation_path) if conversation_path.exists() else []
        tool_calls = [f"{call['function']}({json.dumps(call['args'], ensure_ascii=False)})"
                      for message in messages if message.get("role") == "assistant"
                      for call in (message.get("tool_calls") or [])]
        final_answers = [text_of(m.get("content")) for m in messages
                         if m.get("role") == "assistant" and not m.get("tool_calls")]
        usage = result.get("actor_usage") or {}
        counts = result.get("counts") or {}
        rows.append({
            "Model": entry["label"],
            "Task": result["task"],
            "Lifecycle": result["lifecycle"],
            "Utility": result["utility"],
            "Actor requests": counts.get("actor_requests"),
            "Tool calls proposed": counts.get("tool_proposals"),
            "Tool calls executed": counts.get("dispatches"),
            "Prompt tokens": usage.get("prompt_tokens"),
            "Completion tokens": usage.get("completion_tokens"),
            "Total tokens": usage.get("total_tokens"),
            "Episode seconds": (result.get("timings_seconds") or {}).get("episode_total"),
            "Tool calls made": "\n".join(tool_calls),
            "Final answer": final_answers[-1] if final_answers else "",
            "Backend errors": ", ".join(result.get("backend_errors") or []),
            "Pair ID": result["pair_id"][:12],
            "Attempt ID": result["attempt_id"],
            "Run folder": run_dir.name,
        })
    return rows


def percent(value):
    return None if value is None else round(100 * value, 1)


def summary_row(entry: dict, episodes: list[dict]) -> dict:
    """One row per model. Models without a run keep every metric empty."""
    row = {
        "Model": entry["label"],
        "Source": entry["source"],
        "Hugging Face repo / endpoint model": entry["model_id"],
        "Frozen team config": "yes" if entry["frozen"] else "no (only the actor changed)",
        "Quantization": entry.get("quantization") or "",
        "vLLM version": entry.get("vllm_version") or "",
        "Tool-call check": entry.get("tool_call_check") or "",
        "Status": entry["status"],
        "Episodes planned": None, "Completed": None, "Infrastructure errors": None,
        "Utility (%)": None, "Actor requests": None, "Tool calls proposed": None,
        "Tool calls executed": None, "Total tokens": None, "Mean episode seconds": None,
        "Code identity": "", "Problem": entry.get("error") or "",
    }
    if not entry.get("run_dir"):
        return row
    run_dir = Path(entry["run_dir"])
    metrics = read_json(run_dir / "aggregates" / "summary.json")["metrics"]
    if len(metrics) != 1:
        raise ValueError(f"Expected one actor/mode group in {run_dir}, found {len(metrics)}")
    metric = metrics[0]

    def total(column):
        values = [e[column] for e in episodes if e[column] is not None]
        return sum(values) if values else None

    row.update({
        "Episodes planned": metric["planned"],
        "Completed": metric["completed"],
        "Infrastructure errors": metric["infrastructure_error"],
        "Utility (%)": percent(metric["clean_utility_value"]),
        "Actor requests": total("Actor requests"),
        "Tool calls proposed": total("Tool calls proposed"),
        "Tool calls executed": metric["dispatches"],
        "Total tokens": metric["reported_total_tokens"],
        "Mean episode seconds": None if metric["mean_episode_seconds"] is None
        else round(metric["mean_episode_seconds"], 1),
        "Code identity": read_json(run_dir / "plan.json")["identity"]["implementation"][:12],
    })
    return row


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path: Path, sheets: dict[str, list[dict]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row[h] for h in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for column, header in enumerate(headers, start=1):
            longest = max(len(str(header)), *(len(str(r[header] or "")) for r in rows))
            sheet.column_dimensions[get_column_letter(column)].width = min(longest + 2, 70)
        for row_cells in sheet.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def bar_chart(axis, labels, values, colors, unit, value_format, whole_numbers=False):
    """Bars with the value printed on each; missing values are labeled n/a, not drawn as zero."""
    from matplotlib.ticker import MaxNLocator

    axis.yaxis.set_major_locator(MaxNLocator(integer=whole_numbers))
    positions = range(len(labels))
    bars = axis.bar(positions, [0 if v is None else v for v in values], width=0.6, color=colors, zorder=3)
    top = max([v for v in values if v is not None] or [0]) or 1  # all-zero or all-missing still gets an axis
    for bar, value in zip(bars, values):
        axis.annotate("n/a" if value is None else value_format.format(value) + unit,
                      (bar.get_x() + bar.get_width() / 2, bar.get_height()), xytext=(0, 3),
                      textcoords="offset points", ha="center", fontsize=8.5, color=MUTED_TEXT)
    axis.set_xticks(list(positions), labels, fontsize=8.5, color=TEXT, rotation=20, ha="right")
    axis.set_ylim(0, top * 1.2)
    axis.grid(axis="y", color=GRID, linewidth=1, zorder=0)
    axis.tick_params(length=0, colors=MUTED_TEXT)
    for side in ("top", "right", "left"):
        axis.spines[side].set_visible(False)
    axis.spines["bottom"].set_color(GRID)


def titled_figure(width, height, title, subtitle, columns=1):
    import matplotlib.pyplot as plt

    figure, axes = plt.subplots(1, columns, figsize=(width, height), facecolor=SURFACE, squeeze=False)
    for axis in axes[0]:
        axis.set_facecolor(SURFACE)
    figure.suptitle(title, x=0.02, ha="left", fontsize=13, fontweight="bold", color=TEXT)
    figure.text(0.02, 0.885, subtitle, fontsize=9, color=MUTED_TEXT)
    return figure, axes[0]


def write_charts(directory: Path, summary: list[dict], tasks: str) -> list[Path]:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    directory.mkdir(parents=True, exist_ok=True)
    labels = [row["Model"] for row in summary]
    colors = [SERIES_COLORS[i % len(SERIES_COLORS)] for i in range(len(summary))]
    width = max(7.5, 1.6 * len(labels) + 2)
    paths = []

    figure, (axis,) = titled_figure(width, 4.6, "1. Task utility without protection",
                                    f"Share of episodes where the model completed the user's task ({tasks}). "
                                    "Higher is better; n/a = no run.")
    bar_chart(axis, labels, [row["Utility (%)"] for row in summary], colors, "%", "{:.0f}")
    axis.set_ylim(0, 110)
    axis.set_yticks(range(0, 101, 25), [f"{tick}%" for tick in range(0, 101, 25)])
    figure.subplots_adjust(top=0.8, bottom=0.24)
    paths.append(directory / "1_task_utility.png")

    figure2, axes = titled_figure(width * 1.6, 4.6, "2. What each model did",
                                  "Totals over the model's episodes. Proposed tool calls that were not executed "
                                  "indicate invalid or rejected calls.", columns=3)
    for axis, column in zip(axes, ("Actor requests", "Tool calls proposed", "Tool calls executed")):
        bar_chart(axis, labels, [row[column] for row in summary], colors, "", "{:.0f}", whole_numbers=True)
        axis.set_title(column, loc="left", fontsize=10, color=MUTED_TEXT)
    figure2.subplots_adjust(top=0.74, bottom=0.26, wspace=0.25)
    paths.append(directory / "2_model_activity.png")

    figure3, axes = titled_figure(width * 1.15, 4.6, "3. Cost per model",
                                  "Tokens reported by the model server, and mean episode duration "
                                  "(depends on GPU and quantization).", columns=2)
    bar_chart(axes[0], labels, [row["Total tokens"] for row in summary], colors, "", "{:,.0f}")
    axes[0].set_title("Total tokens", loc="left", fontsize=10, color=MUTED_TEXT)
    bar_chart(axes[1], labels, [row["Mean episode seconds"] for row in summary], colors, " s", "{:.1f}")
    axes[1].set_title("Mean episode duration", loc="left", fontsize=10, color=MUTED_TEXT)
    figure3.subplots_adjust(top=0.74, bottom=0.26, wspace=0.25)
    paths.append(directory / "3_run_cost.png")

    for figure_, path in zip((figure, figure2, figure3), paths):
        figure_.savefig(path, dpi=160, facecolor=SURFACE)
        plt.close(figure_)
    return paths


def build_report(registry_path: Path, output: Path) -> dict[str, Path]:
    registry = read_json(registry_path)
    output.mkdir(parents=True, exist_ok=True)
    episodes, summary = [], []
    for entry in registry["models"]:
        rows = episode_rows(entry) if entry.get("run_dir") else []
        episodes.extend(rows)
        summary.append(summary_row(entry, rows))
    if not summary:
        raise ValueError("The registry lists no models")
    tasks = ", ".join(registry["run_info"].get("tasks", []))
    run_info = [{"Field": key, "Value": json.dumps(value) if isinstance(value, (list, dict)) else value}
                for key, value in registry["run_info"].items()]
    written = {
        "summary.csv": output / "summary.csv",
        "unprotected_baseline_results.xlsx": output / "unprotected_baseline_results.xlsx",
    }
    write_csv(written["summary.csv"], summary)
    if episodes:
        written["episodes.csv"] = output / "episodes.csv"
        write_csv(written["episodes.csv"], episodes)
    sheets = {"Summary": summary, "Run info": run_info}
    if episodes:
        sheets = {"Summary": summary, "Episodes": episodes, "Run info": run_info}
    write_workbook(written["unprotected_baseline_results.xlsx"], sheets)
    for path in write_charts(output / "charts", summary, tasks):
        written[path.name] = path
    return written


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the Person 2 unprotected-baseline report")
    parser.add_argument("--registry", type=Path, required=True, help="registry.json written by the notebook")
    parser.add_argument("--output", type=Path, required=True, help="folder for the workbook, CSVs and charts")
    args = parser.parse_args()
    for name, path in build_report(args.registry, args.output).items():
        print(f"wrote {name}: {path}")


if __name__ == "__main__":
    main()
