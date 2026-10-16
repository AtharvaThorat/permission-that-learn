"""End-to-end test of phase 6 against a fake vLLM server (tests/fake_vllm_server.py), so no GPU is needed.

Checks that the runner starts the server, runs tests in parallel without mixing up their logs,
writes the result files, and reads cached results on a re-run without starting the server again.
The fake model doesn't really solve tasks, so the scores themselves are not checked.
"""

import csv
import json
import socket
import sys
from pathlib import Path

import pytest

from agentdojo.task_suite.load_suites import get_suite
from agentdojo_benchmark import config, local_benchmark

FAKE_SERVER = Path(__file__).with_name("fake_vllm_server.py")
PARALLEL_TESTS = 4


def free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def run_benchmark(monkeypatch, tmp_path: Path, vllm_command: str) -> None:
    models_dir = tmp_path / "models"
    (models_dir / "qwen3").mkdir(parents=True, exist_ok=True)
    (models_dir / "qwen3" / ".download_complete").write_text("fake")  # skip the real download
    monkeypatch.setattr(sys, "argv", [
        "phase6", "--models", "qwen3", "--attacked-tests", "8", "--workers", str(PARALLEL_TESTS),
        "--models-dir", str(models_dir), "--output-dir", str(tmp_path / "output"),
        "--port", str(free_port()), "--vllm-command", vllm_command,
    ])
    local_benchmark.main()


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


@pytest.fixture(scope="module")
def benchmark_output(tmp_path_factory) -> Path:
    """Run the benchmark once against the fake server; the tests below inspect its output."""
    tmp_path = tmp_path_factory.mktemp("phase6")
    monkeypatch = pytest.MonkeyPatch()
    fake_vllm = f'"{Path(sys.executable).as_posix()}" "{FAKE_SERVER.as_posix()}"'
    try:
        run_benchmark(monkeypatch, tmp_path, fake_vllm)
    finally:
        monkeypatch.undo()
    return tmp_path


def test_every_planned_test_has_a_result_and_none_crashed(benchmark_output):
    rows = read_csv(benchmark_output / "output" / "results" / "test_results.csv")
    planned = json.loads((benchmark_output / "output" / "test_plan.json").read_text())
    assert len(rows) == len(planned)
    assert sum(1 for test in planned if test["attack"]) == 8
    assert all(row["error"] == "" for row in rows)


def test_summary_has_one_row_for_the_model_with_numeric_percentages(benchmark_output):
    summary = read_csv(benchmark_output / "output" / "results" / "model_summary.csv")
    assert [row["model"] for row in summary] == ["qwen3-8b-awq"]
    for column in ("baseline_utility_pct", "utility_under_attack_pct", "attack_success_rate_pct"):
        assert 0.0 <= float(summary[0][column]) <= 100.0


def test_breakdown_covers_every_attack_style_and_suite(benchmark_output):
    breakdown = read_csv(benchmark_output / "output" / "results" / "attack_success_breakdown.csv")
    planned = json.loads((benchmark_output / "output" / "test_plan.json").read_text())
    attacked = [test for test in planned if test["attack"]]
    assert {row["group"] for row in breakdown if row["group_by"] == "attack"} == {t["attack"] for t in attacked}
    assert {row["group"] for row in breakdown if row["group_by"] == "suite"} == {t["suite"] for t in attacked}


def test_workbook_has_the_three_result_sheets(benchmark_output):
    from openpyxl import load_workbook

    workbook = load_workbook(benchmark_output / "output" / "results" / "benchmark_results.xlsx")
    assert workbook.sheetnames == ["Model summary", "Attack success breakdown", "Test results"]
    planned = json.loads((benchmark_output / "output" / "test_plan.json").read_text())
    assert workbook["Test results"].max_row == len(planned) + 1  # + header row


def test_tests_really_ran_in_parallel(benchmark_output):
    server_log = (benchmark_output / "output" / "server_logs" / "qwen3.log").read_text()
    most_at_once = max(int(line.split("=")[1]) for line in server_log.splitlines() if "most_requests_at_once" in line)
    assert most_at_once > 1


def test_each_log_belongs_to_its_own_test(benchmark_output):
    """Parallel tests must not write into each other's logs (AgentDojo needs a per-thread fix for this)."""
    log_files = list((benchmark_output / "output" / "logs").rglob("*.json"))
    assert log_files
    for path in log_files:
        log = json.loads(path.read_text())
        suite = get_suite(config.BENCHMARK_VERSION, log["suite_name"])
        user_request = "".join(block["content"] for block in log["messages"][1]["content"])
        assert user_request == suite.get_user_task_by_id(log["user_task_id"]).PROMPT, path
        assert path.parts[-3] == log["user_task_id"]
        assert path.stem == (log["injection_task_id"] or "none")


def test_report_builds_three_charts_and_a_four_page_pdf(benchmark_output):
    import matplotlib

    matplotlib.use("Agg")
    from agentdojo_benchmark.report import build_report, outcome_counts, load_results

    results_dir = benchmark_output / "output" / "results"
    written = build_report(results_dir)
    assert sorted(written) == ["1_model_comparison.png", "2_attack_outcomes.png", "3_attack_success_heatmap.png",
                               "benchmark_report.pdf"]
    assert all(path.stat().st_size > 10_000 for path in written.values())
    pdf_bytes = written["benchmark_report.pdf"].read_bytes()
    assert pdf_bytes.count(b"/Type /Page") - pdf_bytes.count(b"/Type /Pages") == 4  # summary page + 3 charts
    # Every attacked test is counted in exactly one outcome
    counts = outcome_counts(load_results(results_dir))
    assert int(counts.sum(axis=1).iloc[0]) == 8


def test_rerun_uses_cached_results_without_starting_a_server(benchmark_output, monkeypatch):
    # A vLLM command that fails immediately: if the runner tried to start a server, the model would be skipped.
    failing_vllm = f'"{Path(sys.executable).as_posix()}" -c "raise SystemExit(1)"'
    run_benchmark(monkeypatch, benchmark_output, failing_vllm)
    rows = read_csv(benchmark_output / "output" / "results" / "test_results.csv")
    planned = json.loads((benchmark_output / "output" / "test_plan.json").read_text())
    assert len(rows) == len(planned)
    assert all(row["error"] == "" for row in rows)
